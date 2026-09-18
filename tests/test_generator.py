"""Banco de pruebas del generator (docs/05 §8): los 7 casos se generan, marca sintética en cada página y en el
EXIF, ground truth coherente con el modelo (AETOTAL recalculado con `engine/calculo.py`), huellas, trampas de
docs/05 §4.3 presentes, determinismo y coincidencia con lo commiteado en `expedientes/`."""

from __future__ import annotations

import hashlib
import io
import json
import re
import shutil
import subprocess
from datetime import datetime
from decimal import Decimal, localcontext
from pathlib import Path

import openpyxl
import pdfplumber
import pytest
from PIL import Image

from engine.calculo import PRECISION_DECIMAL, calcular
from generator import generar
from generator.casos import caso as caso_modelo
from generator.casos import todos_los_casos
from generator.documentos.escaneo import imagen_ficha_variador
from generator.documentos.imagenes import foto_irrelevante, foto_motor, foto_placa
from generator.documentos.registro import FORMATO_CANONICO
from generator.marcas import MARCA, leer_descripcion_exif
from generator.modelo_caso import Caso

RAIZ = Path(__file__).resolve().parents[1]
EXPEDIENTES = RAIZ / "expedientes"

FICHEROS_ESPERADOS = {"A": 11, "B": 10, "C": 11, "D": 11, "E": 16, "F": 21, "G": 13}
VEREDICTOS = {
    "A": "PREVALIDADO",
    "B": "SUBSANABLE",
    "C": "BLOQUEADO",
    "D": "NO_ELEGIBLE",
    "E": "PREVALIDADO",
    "F": "PREVALIDADO",
    "G": "PREVALIDADO",
}
# docs/05 §2.1, columna "Reglas que deben dispararse"
FALLADAS = {
    "A": [],
    "B": ["R-EVD-04", "R-DOC-01"],
    "C": ["R-CON-01"],
    "D": ["R-AMB-01"],
    "E": [],
    "F": [],
    "G": [],
}
NO_EVALUABLES = {"B": ["R-EVD-01", "R-EVD-02", "R-EVD-03", "R-CON-03"], "C": ["R-CAL-03", "R-CON-06"]}
AETOTAL_A = "305829.6"


@pytest.fixture(scope="module")
def salida(tmp_path_factory) -> Path:
    return tmp_path_factory.mktemp("expedientes")


@pytest.fixture(scope="module")
def generados(salida: Path) -> dict[str, generar.CasoGenerado]:
    return {g.caso.id: g for g in generar.generar_todos(salida)}


@pytest.fixture(scope="module")
def ground_truths(salida: Path, generados) -> dict[str, dict]:
    carpeta = salida / generar.CARPETA_RESULTADOS
    return {
        g.caso.id: json.loads((carpeta / f"{g.caso.carpeta}.json").read_text("utf-8"))
        for g in generados.values()
    }


def _pdfs(carpeta: Path) -> list[Path]:
    return sorted(carpeta.glob("*.pdf"))


def _texto_pdf(ruta: Path) -> str:
    with pdfplumber.open(ruta) as pdf:
        return "\n".join(p.extract_text() or "" for p in pdf.pages)


def _tablas_pdf(ruta: Path) -> list[list[list[str]]]:
    with pdfplumber.open(ruta) as pdf:
        return [t for p in pdf.pages for t in p.extract_tables()]


def _celda(tablas: list[list[list[str]]], etiqueta: str) -> list[str]:
    return [
        fila[1]
        for t in tablas
        for fila in t
        if fila and fila[0] and fila[0].replace("\n", " ").startswith(etiqueta)
    ]


# ---------------------------------------------------------------------------
# Estructura: 7 casos, ficheros, ground truth
# ---------------------------------------------------------------------------


def test_siete_casos_con_sus_ficheros(salida: Path, generados) -> None:
    assert sorted(generados) == list("ABCDEFG")
    for id, g in generados.items():
        carpeta = salida / g.caso.carpeta
        assert carpeta.is_dir()
        ficheros = sorted(p.name for p in carpeta.iterdir())
        assert len(ficheros) == FICHEROS_ESPERADOS[id], (id, ficheros)
        assert ficheros == sorted(f.nombre for f in g.ficheros)
        assert (salida / generar.CARPETA_RESULTADOS / f"{g.caso.carpeta}.json").is_file()


def test_nombres_y_variaciones_de_docs_05(generados) -> None:
    nombres = {id: sorted(f.nombre for f in g.ficheros) for id, g in generados.items()}
    assert not any("registro_funcionamiento" in n for n in nombres["B"])
    assert "06_registro_funcionamiento_MTR-SYN-0001.xlsx" in nombres["A"]
    assert nombres["G"] == sorted(
        ["doc1.pdf", "doc2.pdf", "doc3.pdf", "doc4.pdf", "doc5.pdf", "doc6.pdf", "scan.pdf", "datos.xlsx"]
        + ["foto1.jpg", "foto2.jpg", "foto3.jpg", "foto4.jpg", "notas.pdf"]
    )
    assert sum("06_registro_funcionamiento" in n for n in nombres["F"]) == 3
    assert {n for n in nombres["E"] if "MTR-SYN-0002" in n or "VSD-SYN-0002" in n} == {
        "06_registro_funcionamiento_MTR-SYN-0002.xlsx",
        "07_registro_horas_previo_MTR-SYN-0002.pdf",
        "08_ficha_tecnica_motor_MTR-SYN-0002.pdf",
        "09_ficha_tecnica_variador_VSD-SYN-0002.pdf",
        "10_ficha_tecnica_equipo_accionado_MTR-SYN-0002.pdf",
    }


def test_ground_truth_veredictos_y_reglas(ground_truths) -> None:
    for id, gt in ground_truths.items():
        assert gt["veredicto_esperado"] == VEREDICTOS[id]
        assert gt["reglas_falladas_esperadas"] == FALLADAS[id]
        assert gt["reglas_no_evaluables_esperadas"] == NO_EVALUABLES.get(id, [])
        assert gt["ficha"] == "IND240" and gt["version_ficha"] == "1.1" and gt["version_spec"] == "0.1.0"
        assert gt["fecha_evaluacion"] == "2026-09-18"
        assert gt["marca"] == MARCA


def test_ground_truth_aetotal(ground_truths) -> None:
    for id in ("A", "G"):
        assert ground_truths[id]["aetotal_esperado"] == {
            "exacto": AETOTAL_A,
            "cae": 305829,
            "provisional": False,
        }
    assert ground_truths["B"]["aetotal_esperado"] == {"exacto": AETOTAL_A, "cae": 305829, "provisional": True}
    for id in ("C", "D"):
        a = ground_truths[id]["aetotal_esperado"]
        assert a["exacto"] is None and a["cae"] is None and a["provisional"] is False
    e, f = ground_truths["E"]["aetotal_esperado"], ground_truths["F"]["aetotal_esperado"]
    assert Decimal(e["exacto"]) > Decimal(AETOTAL_A) and e["cae"] == int(
        Decimal(e["exacto"]).to_integral_value("ROUND_DOWN")
    )
    assert Decimal(f["exacto"]) > Decimal(e["exacto"]) and f["cae"] == int(
        Decimal(f["exacto"]).to_integral_value("ROUND_DOWN")
    )
    assert len(ground_truths["E"]["motores"]) == 2 and len(ground_truths["F"]["motores"]) == 3


def test_ground_truth_coherente_con_el_modelo(ground_truths, spec_ind240) -> None:
    """Recalcular con engine/calculo.py desde `motores[]` reproduce `aetotal_esperado.exacto` (docs/05 §8)."""
    from generator.ground_truth import motores_como_unidades

    for id, gt in ground_truths.items():
        unidades = {
            s: {k: Decimal(v) for k, v in vals.items()} for s, vals in motores_como_unidades(gt).items()
        }
        resultado = calcular(spec_ind240.datos, unidades, spec_ind240.tablas)
        esperado = gt["aetotal_esperado"]["exacto"] or gt["aetotal_esperado"]["referencia_no_publicada"]
        assert resultado.total == Decimal(esperado), id
        assert sum(Decimal(m["AEM"]) for m in gt["motores"]) == Decimal(esperado), id
        for unidad, m in zip(resultado.por_unidad, gt["motores"], strict=True):
            assert unidad.salida == Decimal(m["AEM"])
            assert (
                unidad.derivadas["h"]
                == Decimal(m["h"])
                == min(Decimal(m["h_antes"]), Decimal(m["h_despues"] or m["h_antes"]))
            )
            assert unidad.derivadas["p"] == Decimal(m["p"]) and unidad.derivadas[
                "perdidas_ref_kw"
            ] == Decimal(m["perdidas_ref_kw"])
            assert m["fuentes"]["p"] == "tabla:REG1781_CUADRO6"
        # el modelo y el ground truth son la misma cosa
        modelo = caso_modelo(id)
        for m, motor in zip(gt["motores"], modelo.motores, strict=True):
            assert (
                Decimal(m["PM"]) == motor.PM and Decimal(m["N1"]) == motor.N1 and Decimal(m["N2"]) == motor.N2
            )
            assert Decimal(m["h_antes"]) == motor.h_antes
            if id != "B":
                assert Decimal(m["h_despues"]) == motor.h_despues


def test_caso_a_exacto(ground_truths) -> None:
    m = ground_truths["A"]["motores"][0]
    assert (m["PM"], m["N1"], m["N2"], m["h_antes"], m["h"], m["perdidas_ref_kw"]) == (
        "110",
        "1485",
        "1188",
        "6000",
        "6000",
        "5.55",
    )
    with localcontext() as ctx:  # el motor calcula con PRECISION_DECIMAL; el test usa la misma
        ctx.prec = PRECISION_DECIMAL
        assert Decimal(m["p"]) == Decimal("5.55") / Decimal("110")
    assert m["AEM"] == AETOTAL_A


def test_potencias_de_e_y_f_tienen_fila_exacta(spec_ind240) -> None:
    """R-CAL-02 debe dar CUMPLE en E y F: 55 y 160 kW están en el cuadro 6 (docs/05 §2.1)."""
    tabla = spec_ind240.tablas["REG1781_CUADRO6"]
    for motor in caso_modelo("F").motores:
        assert tabla.buscar_exacta(motor.PM) is not None, motor.PM


def test_variables_consolidadas_y_fuentes(ground_truths) -> None:
    for id, gt in ground_truths.items():
        for variables in gt["variables_consolidadas"]["motores"].values():
            for nombre in ("PM", "N1", "N2", "h_antes", "h_despues", "P_prom"):
                assert nombre in variables, (id, nombre)
            assert variables["PM"]["fuente_primaria"] == "ficha_tecnica_motor"
            assert {f["tipo"] for f in variables["PM"]["fuentes"]} == {
                "ficha_tecnica_motor",
                "certificado_instalador",
                "ficha_cumplimentada",
            }
            assert variables["PM"]["fuentes_ocr"][0]["tipo"] == "placa_caracteristicas_foto"
            assert variables["h_antes"]["fuentes"][0]["tipo"] == "registro_horas_previo"
            for fuente in variables["PM"]["fuentes"]:
                assert fuente["pagina"] >= 1 and fuente["fichero"]
        actuacion = gt["variables_consolidadas"]["actuacion"]
        assert actuacion["n_motores"]["valor"] == len(gt["motores"])
        assert {f["tipo"] for f in actuacion["titular_nif"]["fuentes"]} == {
            "ficha_cumplimentada",
            "factura",
            "declaracion_responsable",
            "convenio_cae",
        }
    b = ground_truths["B"]["variables_consolidadas"]["motores"]["MTR-SYN-0001"]
    assert (
        b["N2"]["evidencia"] == "declarado" and b["h_despues"]["valor"] is None and b["h_despues"]["ausente"]
    )
    assert {f["tipo"] for f in b["N2"]["fuentes"]} == {"certificado_instalador", "ficha_cumplimentada"}
    a = ground_truths["A"]["variables_consolidadas"]["motores"]["MTR-SYN-0001"]
    assert a["N2"]["evidencia"] == "derivado" and a["N2"]["fuentes"][0]["tipo"] == "registro_funcionamiento"
    assert a["N2"]["interpretacion"] == "INT-03" and a["h_despues"]["interpretacion"] == "INT-04"
    c = ground_truths["C"]
    pm = c["variables_consolidadas"]["motores"]["MTR-SYN-0001"]["PM"]
    assert pm["valor"] is None and pm["conflicto"] is True
    assert len(c["conflictos_esperados"]) == 1
    valores = {e["tipo"]: e["valor"] for e in c["conflictos_esperados"][0]["evidencias"]}
    assert valores == {"ficha_tecnica_motor": "110", "certificado_instalador": "90"}
    assert (
        ground_truths["D"]["variables_consolidadas"]["motores"]["MTR-SYN-0001"]["tipo_equipo_accionado"][
            "valor"
        ]
        == "bomba_desplazamiento_positivo"
    )


def test_sha256_de_los_documentos(salida: Path, generados, ground_truths) -> None:
    for id, gt in ground_truths.items():
        carpeta = salida / generados[id].caso.carpeta
        assert sorted(d["fichero"] for d in gt["documentos"]) == sorted(p.name for p in carpeta.iterdir())
        for d in gt["documentos"]:
            assert hashlib.sha256((carpeta / d["fichero"]).read_bytes()).hexdigest() == d["sha256"], d[
                "fichero"
            ]


def test_g_pdf_combinado_partes_y_avisos(salida: Path, ground_truths) -> None:
    gt = ground_truths["G"]
    doc1 = next(d for d in gt["documentos"] if d["fichero"] == "doc1.pdf")
    assert doc1["tipo"] == "combinado"
    assert [p["tipo"] for p in doc1["partes"]] == [
        "ficha_cumplimentada",
        "declaracion_responsable",
        "convenio_cae",
    ]
    with pdfplumber.open(salida / "EXP001-G_desordenado" / "doc1.pdf") as pdf:
        assert len(pdf.pages) == doc1["partes"][-1]["pagina_fin"]
        for parte, titulo in zip(
            doc1["partes"],
            ("Ficha IND240 cumplimentada", "Declaración responsable", "Convenio CAE"),
            strict=True,
        ):
            assert titulo in (pdf.pages[parte["pagina_inicio"] - 1].extract_text() or "")
    tipos_aviso = {a["tipo"] for a in gt["avisos_esperados"]}
    assert {
        "documento_no_clasificado",
        "pdf_separado",
        "escaneo_ocr",
        "registro_vinculado_por_hash",
    } <= tipos_aviso
    vinculo = next(a for a in gt["avisos_esperados"] if a["tipo"] == "registro_vinculado_por_hash")
    assert vinculo["fichero"] == "datos.xlsx" and vinculo["nombre_declarado"] != "datos.xlsx"
    assert {d["fichero"] for d in gt["documentos"] if d.get("irrelevante")} == {"notas.pdf", "foto4.jpg"}
    assert {d["fichero"]: d["subtipo"] for d in gt["documentos"] if d["formato"] == "imagen"} == {
        "foto1.jpg": "foto_antes",
        "foto2.jpg": "foto_despues",
        "foto3.jpg": "placa",
        "foto4.jpg": "irrelevante",
    }
    for id in "ABCDEF":
        assert ground_truths[id]["avisos_esperados"] == []


# ---------------------------------------------------------------------------
# Marca sintética: cada página, cada foto
# ---------------------------------------------------------------------------


def test_marca_en_cada_pagina_de_cada_pdf_nativo(salida: Path, generados) -> None:
    paginas = 0
    for g in generados.values():
        for ruta in _pdfs(salida / g.caso.carpeta):
            if ruta.name == "scan.pdf":
                continue
            with pdfplumber.open(ruta) as pdf:
                for pagina in pdf.pages:
                    assert MARCA in (pagina.extract_text() or ""), (ruta.name, pagina.page_number)
                    paginas += 1
    assert paginas > 60


def test_escaneo_de_g_sin_capa_de_texto_pero_con_imagen(salida: Path) -> None:
    with pdfplumber.open(salida / "EXP001-G_desordenado" / "scan.pdf") as pdf:
        assert len(pdf.pages) == 1
        assert (pdf.pages[0].extract_text() or "").strip() == ""
        assert len(pdf.pages[0].images) == 1
        assert pdf.pages[0].width > pdf.pages[0].height  # girado 90°: página apaisada


def test_marca_en_exif_de_cada_foto(salida: Path) -> None:
    motor = caso_modelo("A").motores[0]
    en_memoria = {
        "antes": foto_motor(motor, "ANTES"),
        "despues": foto_motor(motor, "DESPUES"),
        "placa": foto_placa(motor),
        "irrelevante": foto_irrelevante(),
        "escaneo": imagen_ficha_variador(motor),
    }
    en_disco = {p.name: p.read_bytes() for p in sorted((salida / "EXP001-G_desordenado").glob("*.jpg"))}
    assert len(en_disco) == 4
    for nombre, datos in {**en_memoria, **en_disco}.items():
        descripcion = leer_descripcion_exif(Image.open(io.BytesIO(datos)))
        assert descripcion.startswith(MARCA), nombre
    assert "foto ANTES" in leer_descripcion_exif(Image.open(io.BytesIO(en_disco["foto1.jpg"])))
    assert "VSD-SYN-0001" in leer_descripcion_exif(Image.open(io.BytesIO(en_disco["foto2.jpg"])))
    assert "placa" in leer_descripcion_exif(Image.open(io.BytesIO(en_disco["foto3.jpg"])))
    assert "MTR-SYN-0001" not in leer_descripcion_exif(Image.open(io.BytesIO(en_disco["foto4.jpg"])))


@pytest.mark.ocr
def test_ocr_lee_la_placa_y_el_escaneo(salida: Path, tmp_path: Path) -> None:
    if shutil.which("tesseract") is None:
        pytest.skip("tesseract no está instalado")
    placa = tmp_path / "placa.png"
    Image.open(salida / "EXP001-G_desordenado" / "foto3.jpg").save(placa)
    texto = subprocess.run(
        ["tesseract", str(placa), "-", "-l", "spa"], capture_output=True, text=True, check=True
    ).stdout
    assert "MTR-SYN-0001" in texto and "110" in texto and "1485" in texto
    escaneo = tmp_path / "scan.png"
    Image.open(io.BytesIO(imagen_ficha_variador(caso_modelo("A").motores[0]))).rotate(-90, expand=True).save(
        escaneo
    )
    texto = subprocess.run(
        ["tesseract", str(escaneo), "-", "-l", "spa"], capture_output=True, text=True, check=True
    ).stdout
    assert "SOLO PRUEBAS" in texto and "3,90" in texto and "VSD-SYN-0001" in texto


# ---------------------------------------------------------------------------
# Registro de funcionamiento: huella, derivaciones, vinculación
# ---------------------------------------------------------------------------


def _canonico_desde_xlsx(ruta: Path) -> tuple[str, list[tuple[datetime, str, int, Decimal]]]:
    """Implementación independiente del formato canónico documentado (lo que hará engine/registro_xlsx.py)."""
    wb = openpyxl.load_workbook(ruta, read_only=True)
    filas = list(wb["registro"].iter_rows(values_only=True))
    assert filas[0] == ("fecha_hora", "estado", "velocidad_rpm", "potencia_kw")
    datos = [(f[0], f[1], int(f[2]), Decimal(str(f[3])).quantize(Decimal("0.1"))) for f in filas[1:]]
    lineas = [f"{t.strftime('%Y-%m-%dT%H:%M:%S')};{e};{v};{p}" for t, e, v, p in datos]
    return hashlib.sha256("\n".join(lineas).encode("utf-8")).hexdigest(), datos


def test_huella_del_registro_coincide_con_certificado_y_metadatos(
    salida: Path, generados, ground_truths
) -> None:
    for id, gt in ground_truths.items():
        if id == "B":
            assert gt["registro"]["ausente"] is True
            continue
        carpeta = salida / generados[id].caso.carpeta
        assert gt["registro"]["formato_canonico"] == FORMATO_CANONICO
        certificado = next(
            p for p in carpeta.iterdir() if p.name in ("05_certificado_instalador.pdf", "doc3.pdf")
        )
        texto_cert = _texto_pdf(certificado)
        huellas_cert = _celda(_tablas_pdf(certificado), "b) SHA-256 del registro")
        assert len(huellas_cert) == len(gt["motores"])
        for serie, r in gt["registro"]["por_motor"].items():
            ruta = carpeta / r["fichero"]
            huella, datos = _canonico_desde_xlsx(ruta)
            assert huella == r["sha256_canonico"], (id, serie)
            assert huella in huellas_cert and huella in texto_cert
            assert gt["hechos_documentales"][f"registro.hash_declarado[{serie}]"] == huella
            meta = dict(openpyxl.load_workbook(ruta, read_only=True)["metadatos"].iter_rows(values_only=True))
            assert meta["sha256_datos_canonicos"] == huella and meta["num_serie_motor"] == serie
            assert meta["marca"] == MARCA
            # derivaciones de la spec (INT-03, INT-04) reproducen exactamente el modelo
            motor = caso_modelo(id).motor(serie)
            marcha = [d for d in datos if d[1] == "MARCHA"]
            assert Decimal(sum(d[2] for d in marcha)) / len(marcha) == motor.N2
            assert sum(d[3] for d in marcha) / len(marcha) == motor.P_prom
            assert Decimal(len(marcha)) * 15 / 60 * 8760 / (Decimal(len(datos)) * 15 / 60) == motor.h_despues
            assert datos[0][0] == datetime(2026, 3, 5) and datos[-1][0] == datetime(2026, 4, 9, 23, 45)
            assert (datos[-1][0] - datos[0][0]).days + 1 >= 30 and gt["registro"]["dias"] == 36
            assert r["nombre_declarado"] in texto_cert


def test_g_registro_se_vincula_por_hash_no_por_nombre(salida: Path, ground_truths) -> None:
    g, a = ground_truths["G"], ground_truths["A"]
    assert g["registro"]["por_motor"]["MTR-SYN-0001"]["fichero"] == "datos.xlsx"
    assert (
        g["registro"]["por_motor"]["MTR-SYN-0001"]["sha256_canonico"]
        == a["registro"]["por_motor"]["MTR-SYN-0001"]["sha256_canonico"]
    )
    assert "datos.xlsx" not in _texto_pdf(salida / "EXP001-G_desordenado" / "doc3.pdf")


# ---------------------------------------------------------------------------
# Trampas deliberadas (docs/05 §4.3)
# ---------------------------------------------------------------------------


def test_trampa_ficha_variador_declara_3_90_kw(salida: Path, generados) -> None:
    for id, g in generados.items():
        carpeta = salida / g.caso.carpeta
        fichas = [p for p in carpeta.glob("09_ficha_tecnica_variador_*.pdf")]
        if id == "G":
            assert not fichas  # en G va como escaneo sin texto (comprobado por OCR)
            continue
        assert len(fichas) == g.caso.n_motores
        for ficha in fichas:
            assert "3,90 kW" in _texto_pdf(ficha)
            assert _celda(_tablas_pdf(ficha), "Pérdidas declaradas por el fabricante") == ["3,90 kW"]
    for m in generados["A"].ground_truth["motores"]:
        assert m["perdidas_ref_kw"] == "5.55" != "3.90"


def test_trampa_factura_menciona_motor_existente(salida: Path, generados) -> None:
    for g in generados.values():
        factura = next(
            p for p in (salida / g.caso.carpeta).iterdir() if p.name in ("03_factura.pdf", "doc2.pdf")
        )
        texto = _texto_pdf(factura)
        assert "motor existente" in texto
        for m in g.caso.motores:
            assert m.num_serie_variador in texto
        assert "Base imponible" in texto and "IVA" in texto and "Total factura" in texto
        assert [l["categoria"] for l in g.ground_truth["hechos_documentales"]["factura.lineas"]] == [
            "variador"
        ] * g.caso.n_motores + ["instalacion"]


def test_trampa_d_declara_ahorro_fuera_de_ambito(salida: Path, ground_truths) -> None:
    carpeta = salida / "EXP001-D_fuera_ambito"
    ficha = _texto_pdf(carpeta / "01_ficha_IND240_cumplimentada.pdf")
    assert "Ahorro anual estimado" in ficha and "305.829 kWh" in ficha
    assert "305.829 kWh" in _texto_pdf(carpeta / "11_convenio_cae.pdf")
    assert "bomba_desplazamiento_positivo" in _texto_pdf(
        carpeta / "10_ficha_tecnica_equipo_accionado_MTR-SYN-0001.pdf"
    )
    assert "bomba_desplazamiento_positivo" in _texto_pdf(carpeta / "05_certificado_instalador.pdf")
    assert ground_truths["D"]["aetotal_esperado"]["exacto"] is None
    assert "Ahorro anual estimado" not in _texto_pdf(
        salida / "EXP001-A_completo" / "01_ficha_IND240_cumplimentada.pdf"
    )


def test_trampa_c_pm_contradictorio(salida: Path) -> None:
    carpeta = salida / "EXP001-C_contradictorio"
    assert _celda(
        _tablas_pdf(carpeta / "05_certificado_instalador.pdf"), "a) Potencia nominal del motor PM"
    ) == ["90 kW"]
    assert _celda(
        _tablas_pdf(carpeta / "08_ficha_tecnica_motor_MTR-SYN-0001.pdf"), "Potencia nominal PM"
    ) == ["110 kW"]
    assert _celda(
        _tablas_pdf(carpeta / "01_ficha_IND240_cumplimentada.pdf"), "Potencia nominal del motor PM"
    ) == ["110 kW"]


def test_trampa_f_menor_h_es_h_despues_en_m3(ground_truths) -> None:
    motores = {m["num_serie_motor"]: m for m in ground_truths["F"]["motores"]}
    m3 = motores["MTR-SYN-0003"]
    assert Decimal(m3["h_despues"]) < Decimal(m3["h_antes"]) and m3["h"] == m3["h_despues"]
    for serie in ("MTR-SYN-0001", "MTR-SYN-0002"):
        assert (
            Decimal(motores[serie]["h_despues"]) > Decimal(motores[serie]["h_antes"])
            and motores[serie]["h"] == motores[serie]["h_antes"]
        )
    assert ground_truths["F"]["trampas"]["menor_h_es_h_despues"] == ["MTR-SYN-0003"]


def test_trampa_b_n2_solo_declarado(salida: Path, ground_truths) -> None:
    carpeta = salida / "EXP001-B_falta_registro"
    texto = _texto_pdf(carpeta / "05_certificado_instalador.pdf")
    assert "1.188 rpm" in texto and "SHA-256" not in texto
    assert "1.188 rpm" in _texto_pdf(carpeta / "01_ficha_IND240_cumplimentada.pdf")
    assert ground_truths["B"]["motores"][0]["h_despues"] is None


def test_placa_solo_en_imagen_y_pies_de_foto_nativos(salida: Path) -> None:
    informe = salida / "EXP001-A_completo" / "04_informe_fotografico.pdf"
    texto = _texto_pdf(informe)
    assert "Foto 1 – ANTES – motor MTR-SYN-0001" in texto
    assert "Foto 2 – DESPUÉS – motor MTR-SYN-0001 con variador VSD-SYN-0001" in texto
    assert "Placa de características" in texto
    assert "110 kW" not in texto and "1.485" not in texto and "1485" not in texto
    with pdfplumber.open(informe) as pdf:
        assert sum(len(p.images) for p in pdf.pages) == 3


# ---------------------------------------------------------------------------
# Tablas antes que texto, datos inventados, determinismo, lo commiteado
# ---------------------------------------------------------------------------


def test_etiqueta_y_valor_en_tablas(salida: Path) -> None:
    carpeta = salida / "EXP001-A_completo"
    tablas = _tablas_pdf(carpeta / "08_ficha_tecnica_motor_MTR-SYN-0001.pdf")
    assert _celda(tablas, "Nº de serie del motor") == ["MTR-SYN-0001"]
    assert _celda(tablas, "Potencia nominal PM") == ["110 kW"] and _celda(tablas, "Velocidad nominal N1") == [
        "1.485 rpm"
    ]
    tablas = _tablas_pdf(carpeta / "07_registro_horas_previo_MTR-SYN-0001.pdf")
    assert _celda(tablas, "Horas anuales de funcionamiento (h_antes)") == ["6.000 h"]
    tablas = _tablas_pdf(carpeta / "11_convenio_cae.pdf")
    assert _celda(tablas, "Ahorro anual de energía final") == ["305.829 kWh"]
    texto = _texto_pdf(carpeta / "11_convenio_cae.pdf")
    for epigrafe in (
        "Identificación de partes",
        "Título descriptivo",
        "Localización (UTM y referencia catastral)",
        "Ahorro anual (kWh)",
        "Tipo de contraprestación",
        "Vida útil",
        "no suscribir otros convenios",
        "Fecha de firma",
    ):
        assert epigrafe in texto, epigrafe
    ficha = _texto_pdf(carpeta / "01_ficha_IND240_cumplimentada.pdf")
    assert "Firmado electrónicamente por Ana Ficticia Sintética" in ficha
    assert _celda(_tablas_pdf(carpeta / "02_declaracion_responsable.pdf"), "Razón social del titular") == [
        "Industrias Sintéticas del Ebro SL"
    ]
    assert _celda(
        _tablas_pdf(carpeta / "01_ficha_IND240_cumplimentada.pdf"), "Propietario inicial del ahorro"
    ) == ["Industrias Sintéticas del Ebro, S.L."]


def test_todo_es_inventado() -> None:
    from generator.modelo_caso import cif_empresa, nif_persona

    for c in todos_los_casos():
        for e in (c.titular, c.instalador, c.convenio.sujeto_delegado):
            assert (
                "Sintétic" in e.razon_social
                or "Ficticio" in e.razon_social
                or "Inventad" in e.razon_social
                or "Imaginari" in e.razon_social
                or "Supuest" in e.razon_social
            )
            assert e.nif == cif_empresa(e.nif[0], int(e.nif[1:8]))  # control válido
            assert e.representante.nif == nif_persona(int(e.representante.nif[:8]))
        for m in c.motores:
            assert "SYN" in m.num_serie_motor and "SYN" in m.num_serie_variador
        assert "Villasintética" in c.localizacion.direccion


def test_modelo_rechaza_float_y_registro_incoherente() -> None:
    from dataclasses import replace

    motor = caso_modelo("A").motores[0]
    with pytest.raises(TypeError):
        replace(motor, PM=110.0)
    with pytest.raises(ValueError):
        replace(motor, h_despues=Decimal("6571"))


def test_determinismo_dos_generaciones_mismos_bytes(generados, spec_ind240) -> None:
    for id in ("A", "G", "F"):
        otra = generar.generar_caso(caso_modelo(id), spec_ind240)
        assert [(f.nombre, f.sha256) for f in otra.ficheros] == [
            (f.nombre, f.sha256) for f in generados[id].ficheros
        ]
        assert otra.ground_truth == generados[id].ground_truth


def test_lo_commiteado_coincide_con_la_regeneracion(salida: Path, generados) -> None:
    """Regenerar no cambia ni los documentos ni el ground truth commiteados (docs/05 §8.2.1)."""
    if not (EXPEDIENTES / "EXP001-A_completo").is_dir():
        pytest.skip("expedientes/ no generado en este clon")
    for g in generados.values():
        carpeta = EXPEDIENTES / g.caso.carpeta
        assert sorted(p.name for p in carpeta.iterdir()) == sorted(f.nombre for f in g.ficheros)
        for f in g.ficheros:
            assert (carpeta / f.nombre).read_bytes() == f.contenido, (
                f"{g.caso.carpeta}/{f.nombre} difiere: regenera y commitea"
            )
        commiteado = json.loads(
            (EXPEDIENTES / generar.CARPETA_RESULTADOS / f"{g.caso.carpeta}.json").read_text("utf-8")
        )
        nuevo = json.loads(
            (salida / generar.CARPETA_RESULTADOS / f"{g.caso.carpeta}.json").read_text("utf-8")
        )
        assert commiteado == nuevo, f"ground truth de {g.caso.id} difiere: cambio de modelo sin ADR"


def test_cli_solo_un_caso(tmp_path: Path) -> None:
    assert generar.main(["--solo", "a", "--salida", str(tmp_path)]) == 0
    assert sorted(p.name for p in tmp_path.iterdir()) == ["EXP001-A_completo", "_resultados_esperados"]


def test_engine_no_importa_generator() -> None:
    """`engine/` no importa de `generator/` (docs/01 §1).

    Se miran importaciones reales, no menciones en docstrings.
    """
    patron = re.compile(
        r"^\s*(?:from\s+(generator|agentes|salida|tests)\b|import\s+(generator|agentes|salida|tests)\b)", re.M
    )
    for ruta in (RAIZ / "engine").glob("*.py"):
        assert not patron.search(ruta.read_text("utf-8")), ruta.name


def test_caso_es_inmutable_y_las_variaciones_son_del_modelo() -> None:
    casos = {c.id: c for c in todos_los_casos()}
    assert casos["B"].variaciones.sin_registro and casos["C"].variaciones.pm_certificado == {
        "MTR-SYN-0001": Decimal("90")
    }
    assert casos["D"].variaciones.ficha_declara_ahorro and casos["G"].variaciones.desordenado
    assert isinstance(casos["A"], Caso)
    with pytest.raises((AttributeError, TypeError)):  # dataclass frozen
        casos["A"].id = "Z"  # type: ignore[misc]
