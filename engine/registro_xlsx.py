"""Lector del registro de funcionamiento (EVD-01): el xlsx del SCADA a valores derivados con huella.

`leer_registro(doc, spec)` devuelve un `RegistroFuncionamiento` con el periodo, los datos canonicos, su
SHA-256 y las **variables derivadas del registro** (las que la ficha declare: N2, P_prom, h_despues).
Todo con `Decimal`; nunca `float`, ni siquiera al leer una celda.

Que decide la spec y que decide el codigo (regla de oro 4: la ficha es configuracion):

- **Que se deriva** sale de `variables[*].derivacion.fuente == <tipo del documento>`; los nombres (`N2`,
  `P_prom`, `h_despues`), su `interpretacion` (INT-03, INT-04) y su `unidad` son datos de la spec.
- **Como se deriva** lo dice la prosa de `derivacion.metodo`, que se resuelve contra el catalogo de metodos
  implementados (`METODOS`) por palabras clave, igual que `calculo.py` hace con el criterio de redondeo
  (INT-06). Un metodo que no case con ninguno **no se inventa**: el registro deja aviso
  `metodo_no_implementado` y esa variable no se deriva.
- El codigo implementa dos metodos, que son los que la ficha usa hoy:

  | Metodo | Se reconoce por | Que calcula |
  |---|---|---|
  | `media_en_marcha` | "media" + "marcha" + magnitud | media de esa columna en las filas
    con `estado = MARCHA` (INT-03) |
  | `extrapolacion_anual` | "8760" + "marcha" | horas en MARCHA x 8760 / horas del periodo (INT-04) |

**Formato canonico** (INT-05; lo que se hashea y lo que el certificado declara). El generator lo fija en
`generator/documentos/registro.py::FORMATO_CANONICO`; aqui se reimplementa **sin importar `generator`**
(dependencias hacia dentro):

    una linea por fila de la hoja de datos, en su orden, sin cabecera, unidas por "\\n" (sin salto final),
    UTF-8:  fecha_hora;estado;velocidad_rpm;potencia_kw
    - fecha_hora: ISO 8601 con segundos y sin zona (`2026-03-05T00:00:00`)
    - estado: en mayusculas (`MARCHA` | `PARO`)
    - velocidad_rpm: entero
    - potencia_kw: exactamente un decimal, con punto (`59.6`, `0.0`)

Heuristicas de lectura declaradas (van a `docs/03` §8):

- **Hoja de datos**: la primera cuya cabecera tenga las cuatro columnas (fecha, estado, velocidad,
  potencia), buscadas por nombre normalizado y no por posicion. Sin ellas, el registro queda ilegible con
  aviso y sin derivados (nunca una excepcion que tumbe el motor).
- **Intervalo**: el salto mas frecuente entre filas consecutivas; si hay saltos distintos, aviso
  `registro_irregular` y se conserva el mas frecuente. El periodo va de la primera fila a la ultima **mas un
  intervalo** (el ultimo intervalo tambien cuenta), que es lo que da 36 dias en el caso A.
- **Integridad**: el fichero se relee por su `ruta` y se comprueba que su SHA-256 sigue siendo el de la
  ingesta; si no, aviso `fichero_alterado` (la huella de la ingesta manda, no se recalcula).
- El formato de exportacion de los registradores **Schneider Altivar** no se soporta: no hay un fichero real
  publico con el que fijarlo y no se inventa su estructura (queda como `TODO` de ingesta en el informe).
"""

from __future__ import annotations

import hashlib
from collections import Counter
from collections.abc import Mapping, Sequence
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Protocol

COLUMNA_FECHA = "fecha_hora"
COLUMNA_ESTADO = "estado"
COLUMNA_VELOCIDAD = "velocidad_rpm"
COLUMNA_POTENCIA = "potencia_kw"
COLUMNAS_CANONICAS = (COLUMNA_FECHA, COLUMNA_ESTADO, COLUMNA_VELOCIDAD, COLUMNA_POTENCIA)
SENALES_COLUMNA: dict[str, tuple[str, ...]] = {
    COLUMNA_FECHA: ("fecha", "timestamp", "instante"),
    COLUMNA_ESTADO: ("estado", "status", "marcha/paro"),
    COLUMNA_VELOCIDAD: ("velocidad", "rpm", "speed"),
    COLUMNA_POTENCIA: ("potencia", "kw", "power"),
}
ESTADO_MARCHA = "MARCHA"
SEPARADOR = ";"
DECIMAL_POTENCIA = Decimal("0.1")
HORAS_ANO = Decimal(8760)
MINUTOS_HORA = Decimal(60)
ROL_MEDIA_VELOCIDAD = "media_velocidad"
ROL_MEDIA_POTENCIA = "media_potencia"
ROL_EXTRAPOLACION = "extrapolacion_horas"
CLAVE_SERIE_MOTOR = "num_serie_motor"
CLAVE_SERIE_VARIADOR = "num_serie_variador"


class DocumentoRegistro(Protocol):
    """Lo que el lector necesita de un documento de ingesta (sin importar `ingesta`)."""

    doc_id: str
    sha256: str
    nombre: str
    tipo: str | None
    exif: dict[str, str]

    @property
    def ruta(self): ...  # Path


@dataclass(frozen=True)
class Metodo:
    """Un metodo de derivacion implementado y las palabras por las que se reconoce en la prosa."""

    rol: str
    obligatorias: frozenset[str]
    columna: str | None  # columna sobre la que opera (None: no usa columna)

    def casa(self, prosa: str) -> bool:
        return all(palabra in prosa for palabra in self.obligatorias)


METODOS: tuple[Metodo, ...] = (
    Metodo(ROL_MEDIA_VELOCIDAD, frozenset({"media", "marcha", "velocidad"}), COLUMNA_VELOCIDAD),
    Metodo(ROL_MEDIA_POTENCIA, frozenset({"media", "marcha", "potencia"}), COLUMNA_POTENCIA),
    Metodo(ROL_EXTRAPOLACION, frozenset({"8760", "marcha"}), None),
)


@dataclass
class RegistroFuncionamiento:
    """Lectura completa de un registro: periodo, datos canonicos, huella y variables derivadas."""

    doc_id: str
    nombre: str
    sha256_fichero: str
    num_serie_motor: str | None
    num_serie_variador: str | None
    inicio: date | None
    fin: date | None
    dias: Decimal | None
    intervalo_min: Decimal | None
    n_filas: int
    n_marcha: int
    datos_canonicos: str
    sha256_datos: str
    medias_marcha: dict[str, Decimal] = field(default_factory=dict)
    horas_marcha: Decimal | None = None
    horas_periodo: Decimal | None = None
    derivados: dict[str, Decimal] = field(default_factory=dict)  # variable de la spec → valor
    interpretaciones: dict[str, str | None] = field(default_factory=dict)
    unidades: dict[str, str | None] = field(default_factory=dict)
    metodos: dict[str, str] = field(default_factory=dict)  # variable → prosa del metodo de la spec
    nombres_por_rol: dict[str, str] = field(default_factory=dict)  # rol → variable de la spec
    avisos: list[str] = field(default_factory=list)
    legible: bool = True

    def avisar(self, aviso: str) -> None:
        if aviso not in self.avisos:
            self.avisos.append(aviso)

    def derivado(self, variable: str) -> Decimal | None:
        return self.derivados.get(variable)

    def _por_rol(self, rol: str) -> Decimal | None:
        variable = self.nombres_por_rol.get(rol)
        return self.derivados.get(variable) if variable else None

    @property
    def n2(self) -> Decimal | None:
        """Velocidad media en MARCHA (la spec la llama `N2`; INT-03)."""
        return self._por_rol(ROL_MEDIA_VELOCIDAD)

    @property
    def p_prom(self) -> Decimal | None:
        """Potencia media en MARCHA (la spec la llama `P_prom`)."""
        return self._por_rol(ROL_MEDIA_POTENCIA)

    @property
    def h_despues(self) -> Decimal | None:
        """Horas anuales extrapoladas del periodo registrado (la spec las llama `h_despues`; INT-04)."""
        return self._por_rol(ROL_EXTRAPOLACION)

    def a_dict(self) -> dict[str, object]:
        return {
            "doc_id": self.doc_id,
            "nombre": self.nombre,
            "num_serie_motor": self.num_serie_motor,
            "num_serie_variador": self.num_serie_variador,
            "inicio": self.inicio.isoformat() if self.inicio else None,
            "fin": self.fin.isoformat() if self.fin else None,
            "dias": str(self.dias) if self.dias is not None else None,
            "n_filas": self.n_filas,
            "n_marcha": self.n_marcha,
            "sha256_datos": self.sha256_datos,
            "derivados": {k: str(v) for k, v in self.derivados.items()},
            "interpretaciones": dict(self.interpretaciones),
            "avisos": list(self.avisos),
        }


# ---------------------------------------------------------------------------
# Normalizacion de celdas
# ---------------------------------------------------------------------------


def _texto(valor: object) -> str:
    return "" if valor is None else str(valor).strip()


def _normalizar(texto: str) -> str:
    return " ".join(_texto(texto).lower().split())


def _decimal(valor: object) -> Decimal | None:
    """Celda → `Decimal` sin pasar por `float` (openpyxl devuelve numeros nativos)."""
    if valor is None or isinstance(valor, bool):
        return None
    if isinstance(valor, Decimal):
        return valor
    texto = _texto(valor).replace(" ", "").replace(" ", "")
    if not texto:
        return None
    if texto.count(",") == 1 and texto.count(".") == 0:
        texto = texto.replace(",", ".")
    else:
        texto = texto.replace(",", "")
    try:
        return Decimal(texto)
    except InvalidOperation:
        return None


def _fecha_hora(valor: object) -> datetime | None:
    if isinstance(valor, datetime):
        return valor
    if isinstance(valor, date):
        return datetime.combine(valor, datetime.min.time())
    texto = _texto(valor)
    if not texto:
        return None
    for formato in (
        "%Y-%m-%dT%H:%M:%S",
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d %H:%M",
        "%d/%m/%Y %H:%M:%S",
        "%d/%m/%Y %H:%M",
        "%Y-%m-%d",
    ):
        try:
            return datetime.strptime(texto, formato)
        except ValueError:
            continue
    return None


def linea_canonica(momento: datetime, estado: str, velocidad: Decimal, potencia: Decimal) -> str:
    """Una fila en formato canonico (INT-05). Reimplementa el formato del generator, sin importarlo."""
    return SEPARADOR.join(
        (
            momento.strftime("%Y-%m-%dT%H:%M:%S"),
            estado.upper(),
            str(int(velocidad)),
            str(potencia.quantize(DECIMAL_POTENCIA)),
        )
    )


# ---------------------------------------------------------------------------
# Lectura del libro
# ---------------------------------------------------------------------------


def _columnas(cabecera: Sequence[object]) -> dict[str, int] | None:
    """Mapa columna canonica → indice, buscado por nombre normalizado (nunca por posicion)."""
    indices: dict[str, int] = {}
    for posicion, celda in enumerate(cabecera):
        nombre = _normalizar(_texto(celda))
        if not nombre:
            continue
        for columna, senales in SENALES_COLUMNA.items():
            if columna in indices:
                continue
            if any(senal in nombre for senal in senales):
                indices[columna] = posicion
                break
    if all(columna in indices for columna in COLUMNAS_CANONICAS):
        return indices
    return None


def _hoja_de_datos(libro):
    for nombre in libro.sheetnames:
        hoja = libro[nombre]
        cabecera = next(hoja.iter_rows(min_row=1, max_row=1, values_only=True), ())
        indices = _columnas(cabecera or ())
        if indices is not None:
            return hoja, indices
    return None, None


def _intervalo(momentos: list[datetime]) -> tuple[Decimal | None, bool]:
    """Salto mas frecuente entre filas, en minutos, y si el registro es regular."""
    if len(momentos) < 2:
        return None, True
    saltos = Counter(
        int((momentos[i + 1] - momentos[i]).total_seconds() // 60) for i in range(len(momentos) - 1)
    )
    (mas_frecuente, veces), *resto = saltos.most_common()
    if mas_frecuente <= 0:
        return None, False
    return Decimal(mas_frecuente), not resto


# ---------------------------------------------------------------------------
# Derivaciones desde la spec
# ---------------------------------------------------------------------------


def derivaciones_de(spec: object, fuente: str | None) -> dict[str, Mapping[str, object]]:
    """Variables de la spec cuya `derivacion.fuente` es este tipo de documento."""
    if spec is None or not fuente:
        return {}
    variables = getattr(spec, "variables", None) or {}
    encontradas: dict[str, Mapping[str, object]] = {}
    for nombre, declaracion in variables.items():
        derivacion = declaracion.get("derivacion") if hasattr(declaracion, "get") else None
        if isinstance(derivacion, Mapping) and _texto(derivacion.get("fuente")) == fuente:
            encontradas[nombre] = declaracion
    return encontradas


def _metodo_de(prosa: str) -> Metodo | None:
    normalizada = _normalizar(prosa)
    for metodo in METODOS:
        if metodo.casa(normalizada):
            return metodo
    return None


def _derivar(registro: RegistroFuncionamiento, spec: object, fuente: str | None) -> None:
    for variable, declaracion in derivaciones_de(spec, fuente).items():
        derivacion = declaracion.get("derivacion") or {}
        prosa = _texto(derivacion.get("metodo"))
        registro.metodos[variable] = prosa
        registro.interpretaciones[variable] = derivacion.get("interpretacion")
        registro.unidades[variable] = declaracion.get("unidad")
        metodo = _metodo_de(prosa)
        if metodo is None:
            registro.avisar(
                f"metodo_no_implementado: '{registro.nombre}' no deriva {variable}: el metodo declarado "
                f"en la ficha ({prosa!r}) no esta implementado; no se inventa un criterio"
            )
            continue
        registro.nombres_por_rol.setdefault(metodo.rol, variable)
        valor: Decimal | None = None
        if metodo.columna is not None:
            valor = registro.medias_marcha.get(metodo.columna)
        elif registro.horas_periodo and registro.horas_marcha is not None:
            valor = registro.horas_marcha * HORAS_ANO / registro.horas_periodo
        if valor is None:
            registro.avisar(f"sin_datos: '{registro.nombre}' no permite derivar {variable}")
            continue
        registro.derivados[variable] = valor


# ---------------------------------------------------------------------------
# Entrada publica
# ---------------------------------------------------------------------------


def leer_registro(doc: DocumentoRegistro, spec: object = None) -> RegistroFuncionamiento:
    """Lee el xlsx del registro. Nunca lanza: un fichero ilegible devuelve un registro `legible = False`."""
    registro = RegistroFuncionamiento(
        doc_id=doc.doc_id,
        nombre=doc.nombre,
        sha256_fichero=doc.sha256,
        num_serie_motor=(doc.exif or {}).get(CLAVE_SERIE_MOTOR) or None,
        num_serie_variador=(doc.exif or {}).get(CLAVE_SERIE_VARIADOR) or None,
        inicio=None,
        fin=None,
        dias=None,
        intervalo_min=None,
        n_filas=0,
        n_marcha=0,
        datos_canonicos="",
        sha256_datos="",
    )
    try:
        datos = doc.ruta.read_bytes()
    except OSError as exc:
        registro.legible = False
        registro.avisar(f"documento_ilegible: '{doc.nombre}' no se pudo releer ({exc.__class__.__name__})")
        return registro
    if hashlib.sha256(datos).hexdigest() != doc.sha256:
        registro.avisar(
            f"fichero_alterado: '{doc.nombre}' ha cambiado despues de la ingesta; manda la huella original"
        )
    try:
        from openpyxl import load_workbook

        libro = load_workbook(BytesIO(datos), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - un xlsx roto no rompe el motor
        registro.legible = False
        registro.avisar(f"documento_ilegible: '{doc.nombre}' no se abre como xlsx ({exc.__class__.__name__})")
        return registro
    try:
        hoja, indices = _hoja_de_datos(libro)
        if hoja is None or indices is None:
            registro.legible = False
            registro.avisar(
                f"registro_sin_columnas: '{doc.nombre}' no tiene las columnas de un registro de "
                f"funcionamiento ({', '.join(COLUMNAS_CANONICAS)})"
            )
            return registro
        _leer_filas(registro, hoja, indices)
    finally:
        libro.close()
    _derivar(registro, spec, getattr(doc, "tipo", None))
    return registro


def _leer_filas(registro: RegistroFuncionamiento, hoja, indices: Mapping[str, int]) -> None:
    lineas: list[str] = []
    momentos: list[datetime] = []
    sumas: dict[str, Decimal] = {COLUMNA_VELOCIDAD: Decimal(0), COLUMNA_POTENCIA: Decimal(0)}
    n_marcha = 0
    n_filas = 0
    descartadas = 0
    for fila in hoja.iter_rows(min_row=2, values_only=True):
        if fila is None or all(celda is None for celda in fila):
            continue
        momento = _fecha_hora(_celda(fila, indices[COLUMNA_FECHA]))
        estado = _texto(_celda(fila, indices[COLUMNA_ESTADO])).upper()
        velocidad = _decimal(_celda(fila, indices[COLUMNA_VELOCIDAD]))
        potencia = _decimal(_celda(fila, indices[COLUMNA_POTENCIA]))
        if momento is None or not estado or velocidad is None or potencia is None:
            descartadas += 1
            continue
        n_filas += 1
        momentos.append(momento)
        lineas.append(linea_canonica(momento, estado, velocidad, potencia))
        if estado == ESTADO_MARCHA:
            n_marcha += 1
            sumas[COLUMNA_VELOCIDAD] += velocidad
            sumas[COLUMNA_POTENCIA] += potencia
    registro.n_filas = n_filas
    registro.n_marcha = n_marcha
    registro.datos_canonicos = "\n".join(lineas)
    registro.sha256_datos = hashlib.sha256(registro.datos_canonicos.encode("utf-8")).hexdigest()
    if descartadas:
        registro.avisar(
            f"filas_descartadas: '{registro.nombre}' tiene {descartadas} filas incompletas que no entran "
            "en los datos canonicos"
        )
    if not n_filas:
        registro.legible = False
        registro.avisar(f"registro_vacio: '{registro.nombre}' no tiene filas legibles")
        return
    if n_marcha:
        for columna, suma in sumas.items():
            registro.medias_marcha[columna] = suma / Decimal(n_marcha)
    else:
        registro.avisar(f"sin_marcha: '{registro.nombre}' no tiene ninguna fila en estado {ESTADO_MARCHA}")
    intervalo, regular = _intervalo(momentos)
    registro.intervalo_min = intervalo
    if not regular:
        registro.avisar(
            f"registro_irregular: '{registro.nombre}' tiene intervalos distintos; se toma el mas frecuente"
        )
    if intervalo is not None:
        registro.horas_periodo = Decimal(n_filas) * intervalo / MINUTOS_HORA
        registro.horas_marcha = Decimal(n_marcha) * intervalo / MINUTOS_HORA
        registro.dias = registro.horas_periodo / Decimal(24)
        registro.inicio = momentos[0].date()
        from datetime import timedelta

        registro.fin = (momentos[-1] + timedelta(minutes=int(intervalo))).date()
    else:
        registro.inicio = momentos[0].date()
        registro.fin = momentos[-1].date()
        registro.avisar(f"sin_intervalo: '{registro.nombre}' no permite deducir el intervalo de muestreo")


def _celda(fila: Sequence[object], indice: int) -> object:
    return fila[indice] if indice < len(fila) else None
